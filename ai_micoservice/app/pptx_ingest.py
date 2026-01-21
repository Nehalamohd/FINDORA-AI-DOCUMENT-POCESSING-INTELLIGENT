"""
PowerPoint (PPTX) ingestion module for extracting text and creating embeddings.
"""
from pptx import Presentation
from app.embedding import embed
from app.database import SessionLocal
from app.models import Document, Page, Chunk, Embedding
from sqlalchemy import text
from app.logger import logger

#Extracts all text slide by slide from a PPTX file
def extract_text_from_pptx(pptx_path: str):
    """
    Extracts all text slide by slide from a PPTX file, including slide notes.
    """
    
    prs = Presentation(pptx_path)
    for i, slide in enumerate(prs.slides):
        text_runs = []
        #loop over all shapes in slide like table ,box
        for shape in slide.shapes:
            if hasattr(shape, "text"): #ensure shape has text property
                text_runs.append(shape.text)
        
        # Also try to get notes from slide
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes:
                text_runs.append(f"\nNotes: {notes}")
                
        yield i + 1, "\n".join(text_runs)

#for extracting text slide by slide and storing in db
def process_pptx_slide(document_id: int, slide_no: int, slide_text: str):
    """Processes a single PPTX slide: stores text and creates embeddings."""
    db = SessionLocal()
    try:
        if not slide_text.strip():
            return None
            
        # 1. Store Page (slide)
        try:
            page = Page(document_id=document_id, page_number=slide_no, content=slide_text)
            db.add(page)
            db.commit()
            db.refresh(page)
            page_id = page.id
            logger.debug(f"Saved slide {slide_no} for document {document_id}")

            # 2. Chunk & Embed
            embeddings = embed([slide_text])
            if embeddings:
                embedding = embeddings[0]
                chunk = Chunk(document_id=document_id, page_id=page_id, chunk_index=0, content=slide_text)
                db.add(chunk)
                db.commit()
                db.refresh(chunk)
                chunk_id = chunk.id
                
                db.execute(
                    text("INSERT INTO embeddings (chunk_id, embedding) VALUES (:chunk_id, CAST(:embedding AS vector))"),
                    {"chunk_id": chunk_id, "embedding": str(embedding.tolist()) if hasattr(embedding, 'tolist') else str(embedding)}
                )
                db.commit()
            return page_id
        except Exception as db_e:
            logger.error(f"Database error saving slide {slide_no} for doc {document_id}: {str(db_e)}")
            db.rollback()
            return None
    except Exception as general_e:
        logger.error(f"System failure processing slide {slide_no} for doc {document_id}: {str(general_e)}")
        return None
    finally:
        db.close()

def ingest_pptx(file_path: str, filename: str, task_id: str = None, flow_id: str = None, username: str = None):
    """
    Main entry point for ingesting a PPTX file. 
    Extracts text, creates embeddings, and updates the database.
    """
    db = SessionLocal()
    try:
        doc = Document(
            filename=filename,
            file_path=file_path,
            file_type="pptx",
            status="processing",
            task_id=task_id,
            flow_id=flow_id
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)
        document_id = doc.id

        logger.info(f"Processing {filename} (PPTX)...")
        
        for slide_no, slide_text in extract_text_from_pptx(file_path):
            process_pptx_slide(document_id, slide_no, slide_text)

        doc.status = 'completed'
        db.commit()
        logger.info(f"Ingestion completed for {filename} (ID: {document_id})")
        return document_id
    except Exception as e:
        logger.error(f"Ingestion failed for {filename}: {str(e)}", exc_info=True)
        try:
            db.rollback()
            db.execute(
                text("UPDATE documents SET status = 'failed', error_message = :err WHERE task_id = :tid"),
                {"err": str(e), "tid": task_id}
            )
            db.commit()
        except Exception as rollback_e:
            logger.error(f"Failed to update document status to 'failed' for {filename}: {str(rollback_e)}")
        raise e
    finally:
        db.close()
